"""Bulk company data import (CSV or Excel) with a preview before anything is saved.

One-file modes:
- Excel workbook with sheets named after each kind (locations, employees, …)
- Freeform CSV/Excel/text: AI maps it into the same kinds, then the same validators run
"""

from __future__ import annotations

import csv
import io
import json
import re
from typing import Any, Callable, Optional

from sqlmodel import Session, select

from app.agent.places import PLACE_TYPES
from app.models.company import CostCentre, Department, KnowledgeEntry, SiteService
from app.models.org import Location, Person, Resource, StockSku, Vendor

TEMPLATES: dict[str, list[str]] = {
    "locations": ["code", "name", "type", "parent_code", "aliases"],
    "employees": ["email", "name", "department", "role", "manager_email", "location"],
    "departments": ["code", "name", "categories", "primary_email", "backup_emails", "approver_email", "sla_hours",
                    "spend_approval_limit", "location"],
    "resources": ["name", "type", "capacity", "floor", "video_conferencing", "display", "near_department"],
    "vendors": ["name", "category", "contact_email", "location"],
    "stock": ["sku_code", "name", "qty_on_hand", "reorder_min", "reorder_qty", "unit", "location", "vendor_email"],
    "knowledge": ["key", "section", "title", "content"],
    "site_services": ["site", "code", "name", "delivery_model", "owner_department", "unit", "rate", "included_limit",
                      "premium_triggers", "premium_rate", "provider"],
    "cost_centres": ["code", "name", "department", "approver_name", "approver_email"],
}

# Parents / places before teams / people that reference them.
IMPORT_ORDER = [
    "locations",
    "departments",
    "employees",
    "resources",
    "vendors",
    "stock",
    "knowledge",
    "site_services",
    "cost_centres",
]

EXAMPLES: dict[str, list[str]] = {
    "site_services": ["Corporate Office", "high_tea", "High tea / snacks (special menu)", "CHARGEABLE", "CAFETERIA",
                      "per_person", "350", "0", "", "0", "In-house cafeteria"],
    "cost_centres": ["SL-1101", "Sales", "Sales", "Rohan Mehta", "rohan.mehta@yourco.com"],
    "locations": ["PUN", "Pune Hinjewadi", "BRANCH", "WEST", "Hinjewadi; Pune office"],
    "employees": ["riya.shah@yourco.com", "Riya Shah", "Engineering", "EMPLOYEE", "priya.manager@yourco.com", "PUN"],
    # roles: EMPLOYEE | COMPANY_ADMIN | BRANCH_ADMIN | TEAM_OPERATOR (mail authority, not dashboard)
    "departments": ["FACILITIES_PUN", "Facilities - Pune", "maintenance; hvac; electrical", "facilities.pune@yourco.com",
                    "", "", "8", "0", "PUN"],
    "resources": ["Board Room 1", "MEETING_ROOM", "12", "3", "yes", "yes", "Sales"],
    "vendors": ["QuickRide Corporate Cabs", "transport", "bookings@quickride.com", "Gurugram"],
    "stock": ["STAPLER", "Stapler", "12", "4", "24", "each", "PUN", "supplies@vendor.com"],
    "knowledge": ["faq_canteen", "faq", "Canteen timings", "Canteen is open 8:30 AM - 6 PM on weekdays."],
}

_EMAIL = re.compile(r"^[\w.+-]+@[\w-]+(?:\.[\w-]+)+$")


def template_csv(kind: str) -> str:
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(TEMPLATES[kind])
    writer.writerow(EXAMPLES[kind])
    return buf.getvalue()


def company_workbook_bytes() -> bytes:
    """Multi-sheet Excel template: one sheet per import kind with header + example row."""
    from openpyxl import Workbook

    wb = Workbook()
    first = True
    for kind in IMPORT_ORDER:
        ws = wb.active if first else wb.create_sheet(kind)
        if first:
            ws.title = kind
            first = False
        ws.append(TEMPLATES[kind])
        ws.append(EXAMPLES[kind])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _sheet_rows(rows: list[tuple]) -> list[dict[str, str]]:
    if not rows:
        return []
    header = [_key(h) for h in rows[0]]
    return [
        {header[i]: ("" if v is None else str(v).strip()) for i, v in enumerate(r) if i < len(header) and header[i]}
        for r in rows[1:]
        if any(v not in (None, "") for v in r)
    ]


def read_rows(filename: str, content: bytes) -> list[dict[str, str]]:
    name = (filename or "").lower()
    if name.endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        return _sheet_rows(list(wb.active.iter_rows(values_only=True)))
    text = content.decode("utf-8-sig", errors="replace")
    reader = csv.DictReader(io.StringIO(text))
    return [{_key(k): (v or "").strip() for k, v in row.items() if k} for row in reader if any((v or "").strip() for v in row.values())]


def read_workbook_bundles(filename: str, content: bytes) -> dict[str, list[dict[str, str]]]:
    """Map sheet names (or a single CSV with a `kind` column) to template kinds."""
    name = (filename or "").lower()
    if name.endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        out: dict[str, list[dict[str, str]]] = {}
        for sheet in wb.worksheets:
            kind = _match_kind(sheet.title)
            if not kind:
                continue
            rows = _sheet_rows(list(sheet.iter_rows(values_only=True)))
            if rows:
                out[kind] = rows
        return out

    rows = read_rows(filename, content)
    if not rows:
        return {}
    if any("kind" in r for r in rows):
        out = {}
        for row in rows:
            kind = _match_kind(row.get("kind") or "")
            if not kind:
                continue
            clean = {k: v for k, v in row.items() if k != "kind"}
            out.setdefault(kind, []).append(clean)
        return out
    # Single-kind file: guess from headers
    headers = set(rows[0].keys())
    for kind, cols in TEMPLATES.items():
        needed = set(cols[:2])  # code+name or email+name etc.
        if needed <= headers:
            return {kind: rows}
    return {}


def _match_kind(label: str) -> Optional[str]:
    raw = _key(label)
    if raw in TEMPLATES:
        return raw
    aliases = {
        "places": "locations",
        "branches": "locations",
        "offices": "locations",
        "people": "employees",
        "staff": "employees",
        "teams": "departments",
        "depts": "departments",
        "rooms": "resources",
        "faqs": "knowledge",
        "policies": "knowledge",
        "services": "site_services",
        "cost_centers": "cost_centres",
        "costcentres": "cost_centres",
    }
    return aliases.get(raw)


def _key(value: Any) -> str:
    return re.sub(r"[^a-z0-9]+", "_", str(value or "").strip().lower()).strip("_")


def _list(value: str) -> list[str]:
    return [p.strip() for p in re.split(r"[;,|]", value or "") if p.strip()]


def _bool(value: str) -> bool:
    return str(value or "").strip().lower() in {"yes", "y", "true", "1", "available"}


def _num(value: str, default: float = 0.0) -> float:
    try:
        return float(str(value).replace(",", "")) if str(value).strip() else default
    except ValueError:
        return default


def _places_cache(session: Session, tid: str) -> list[Location]:
    """One locations query per import pass (location_by_ref used to re-select every row)."""
    cache = getattr(session, "_import_places_cache", None)
    if cache is None or cache[0] != tid:
        rows = list(session.exec(select(Location).where(Location.tenant_id == tid)).all())
        cache = (tid, rows)
        setattr(session, "_import_places_cache", cache)
    return cache[1]


def _invalidate_places_cache(session: Session) -> None:
    if hasattr(session, "_import_places_cache"):
        delattr(session, "_import_places_cache")


def _people_by_email(session: Session, tid: str) -> dict[str, Person]:
    cache = getattr(session, "_import_people_cache", None)
    if cache is None or cache[0] != tid:
        rows = list(session.exec(select(Person).where(Person.tenant_id == tid)).all())
        cache = (tid, {p.email.lower(): p for p in rows if p.email})
        setattr(session, "_import_people_cache", cache)
    return cache[1]


def _depts_by_code(session: Session, tid: str) -> dict[str, Department]:
    cache = getattr(session, "_import_depts_cache", None)
    if cache is None or cache[0] != tid:
        rows = list(session.exec(select(Department).where(Department.tenant_id == tid)).all())
        cache = (tid, {d.code.upper(): d for d in rows if d.code})
        setattr(session, "_import_depts_cache", cache)
    return cache[1]


def _resources_by_name(session: Session, tid: str) -> dict[str, Resource]:
    cache = getattr(session, "_import_resources_cache", None)
    if cache is None or cache[0] != tid:
        rows = list(session.exec(select(Resource).where(Resource.tenant_id == tid)).all())
        cache = (tid, {r.name.lower(): r for r in rows if r.name})
        setattr(session, "_import_resources_cache", cache)
    return cache[1]


def location_by_ref(session: Session, tid: str, ref: str) -> Location | None:
    """A place by its code (preferred) or exact name."""
    ref = (ref or "").strip()
    if not ref:
        return None
    rows = _places_cache(session, tid)
    by_code = [p for p in rows if (p.code or "").lower() == ref.lower()]
    if by_code:
        return by_code[0]
    by_name = [p for p in rows if p.name.lower() == ref.lower()]
    if len(by_name) > 1:
        raise ValueError(f"more than one place is called '{ref}' - use its code")
    return by_name[0] if by_name else None


def _location_or_error(session: Session, tid: str, ref: str) -> Location | None:
    if not (ref or "").strip():
        return None
    place = location_by_ref(session, tid, ref)
    if place is None:
        raise ValueError(f"unknown location '{ref}' - import locations first")
    return place


def _parents_first(rows: list[dict[str, str]]) -> list[tuple[int, dict[str, str]]]:
    """A locations sheet can list children before their parents; import parents first."""
    numbered = list(enumerate(rows, start=2))
    codes = {(r.get("code") or "").strip().lower(): r for _, r in numbered if (r.get("code") or "").strip()}

    def depth(row: dict[str, str]) -> int:
        d, seen = 0, set()
        parent = (row.get("parent_code") or "").strip().lower()
        while parent in codes and parent not in seen:
            seen.add(parent)
            d += 1
            parent = (codes[parent].get("parent_code") or "").strip().lower()
        return d

    return sorted(numbered, key=lambda item: depth(item[1]))


def _import_kind_rows(
    session: Session, tenant_id: str, kind: str, rows: list[dict[str, str]], *, apply: bool
) -> dict:
    handler: Callable = globals()[f"_import_{kind}"]
    created, updated, errors = 0, 0, []
    preview = []
    ordered = _parents_first(rows) if kind == "locations" else list(enumerate(rows, start=2))
    # Locations need frequent flush for parent_code; other kinds batch heavily.
    flush_every = 1 if kind == "locations" else 200
    pending = 0
    for i, row in ordered:
        try:
            action, label = handler(session, tenant_id, row, apply=apply)
        except ValueError as exc:
            errors.append({"row": i, "kind": kind, "error": str(exc)})
            continue
        created += action == "new"
        updated += action == "update"
        if apply:
            pending += 1
            if pending >= flush_every:
                session.flush()
                if kind == "locations":
                    _invalidate_places_cache(session)
                pending = 0
        if len(preview) < 50:
            preview.append({"row": i, "kind": kind, "action": action, "item": label})
    if apply and pending:
        session.flush()
        if kind == "locations":
            _invalidate_places_cache(session)
    preview.sort(key=lambda p: p["row"])
    return {
        "kind": kind,
        "new": created,
        "updated": updated,
        "errors": errors,
        "preview": preview,
        "total": len(rows),
    }


def run_import(session: Session, tenant_id: str, kind: str, rows: list[dict[str, str]], *, commit: bool) -> dict:
    if kind not in TEMPLATES:
        raise ValueError(f"Unknown import type: {kind}")
    # Locations must exist during a preview too, so later rows can find their parent.
    apply = commit or kind == "locations"
    result = _import_kind_rows(session, tenant_id, kind, rows, apply=apply)
    if commit:
        session.commit()
    else:
        session.rollback()
    return {**result, "committed": commit}


def run_bundle_import(
    session: Session, tenant_id: str, bundles: dict[str, list[dict[str, str]]], *, commit: bool, source: str = "structured"
) -> dict:
    """Import several kinds in dependency order.

    Preview (commit=False) only applies locations (so place refs resolve) and dry-runs
    the rest — a full apply+rollback of 1k+ rows freezes Render's single worker.

    Commit mode saves each sheet as it finishes so a mid-file timeout doesn't lose
    earlier sheets, and raises the session statement timeout for bulk loads.
    """
    from sqlalchemy import text

    parts: list[dict] = []
    if commit:
        session.execute(text("SET LOCAL statement_timeout = '300s'"))
    for kind in IMPORT_ORDER:
        rows = bundles.get(kind) or []
        if not rows:
            continue
        apply = commit or kind == "locations"
        part = _import_kind_rows(session, tenant_id, kind, rows, apply=apply)
        parts.append(part)
        if commit:
            session.commit()
            _invalidate_places_cache(session)
            session.execute(text("SET LOCAL statement_timeout = '300s'"))
    if not commit:
        session.rollback()
    errors = [e for p in parts for e in p["errors"]]
    preview = [row for p in parts for row in p["preview"]][:80]
    return {
        "mode": "company",
        "source": source,
        "committed": commit,
        "parts": parts,
        "kinds": [p["kind"] for p in parts],
        "new": sum(p["new"] for p in parts),
        "updated": sum(p["updated"] for p in parts),
        "errors": errors,
        "preview": preview,
        "total": sum(p["total"] for p in parts),
    }


def _file_text_for_ai(filename: str, content: bytes) -> str:
    name = (filename or "").lower()
    if name.endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        chunks: list[str] = []
        for sheet in wb.worksheets:
            rows = list(sheet.iter_rows(values_only=True))
            if not rows:
                continue
            buf = io.StringIO()
            writer = csv.writer(buf)
            for r in rows:
                writer.writerow(["" if v is None else v for v in r])
            chunks.append(f"### sheet: {sheet.title}\n{buf.getvalue().strip()}")
        return "\n\n".join(chunks)
    return content.decode("utf-8-sig", errors="replace")


def _normalize_bundles(raw: dict) -> dict[str, list[dict[str, str]]]:
    out: dict[str, list[dict[str, str]]] = {}
    for key, rows in (raw or {}).items():
        kind = _match_kind(str(key))
        if not kind or not isinstance(rows, list):
            continue
        cols = TEMPLATES[kind]
        cleaned: list[dict[str, str]] = []
        for row in rows:
            if not isinstance(row, dict):
                continue
            item = {c: str(row.get(c, "") if row.get(c) is not None else "").strip() for c in cols}
            # Also accept close aliases from the model
            for src, dst in list(row.items()):
                k = _key(src)
                if k in cols and not item.get(k):
                    item[k] = "" if dst is None else str(dst).strip()
            if any(item.values()):
                cleaned.append(item)
        if cleaned:
            out[kind] = cleaned
    return out


def structure_freeform_with_ai(filename: str, content: bytes) -> dict[str, list[dict[str, str]]]:
    """Ask the AI to map a messy org dump into our template kinds. Validators still own the write."""
    from app.ai.gemini import GeminiProvider, HeuristicProvider
    from app.core.config import get_settings

    text = _file_text_for_ai(filename, content).strip()
    if not text:
        raise ValueError("The file has no readable text")

    # Already-structured JSON dump
    if text.lstrip().startswith("{"):
        try:
            parsed = json.loads(text)
            bundles = _normalize_bundles(parsed if isinstance(parsed, dict) else {})
            if bundles:
                return bundles
        except json.JSONDecodeError:
            pass

    schema_hint = {kind: TEMPLATES[kind] for kind in IMPORT_ORDER}
    system = (
        "You convert company org / admin desk data into a JSON object for database import. "
        "Return ONLY JSON with optional keys: "
        + ", ".join(IMPORT_ORDER)
        + ". Each key maps to an array of objects. Use exactly these fields per key: "
        + json.dumps(schema_hint)
        + ". Use place codes in location / parent_code / department location fields. "
        "Infer BRANCH/CITY/REGION types when obvious. Leave unknown fields as empty strings. "
        "Do not invent employee emails. Categories are semicolon-separated lowercase tokens."
    )
    payload = {"filename": filename or "upload", "raw_text": text[:120_000]}

    settings = get_settings()
    if settings.gemini_api_key or settings.gemini_api_key_secondary:
        data, _meta = GeminiProvider().generate_json(system_instruction=system, payload=payload, temperature=0.1)
    else:
        data, _meta = HeuristicProvider().generate_json(system_instruction=system, payload=payload, temperature=0.1)

    bundles = _normalize_bundles(data)
    if not bundles:
        raise ValueError(
            "Could not map this file into company data. Use the multi-sheet company template, "
            "or a clearer dump of branches, teams and employees."
        )
    return bundles


def prepare_company_import(filename: str, content: bytes) -> tuple[dict[str, list[dict[str, str]]], str]:
    """Structured workbook/CSV first; AI only when sheets/headers are not already in template shape."""
    name = (filename or "").lower()
    head = content.lstrip()[:1]
    if name.endswith(".json") or head == b"{":
        return structure_freeform_with_ai(filename, content), "ai"
    bundles = read_workbook_bundles(filename, content)
    if bundles:
        source = "workbook" if name.endswith((".xlsx", ".xlsm")) else "structured"
        return bundles, source
    return structure_freeform_with_ai(filename, content), "ai"


def _import_locations(session: Session, tid: str, row: dict, *, apply: bool) -> tuple[str, str]:
    code = (row.get("code") or "").strip().upper()
    name = (row.get("name") or "").strip()
    ptype = (row.get("type") or "BRANCH").strip().upper()
    if not code or not name:
        raise ValueError("code and name are required")
    if ptype not in PLACE_TYPES:
        raise ValueError(f"type must be one of {', '.join(sorted(PLACE_TYPES))} (got '{ptype}')")
    parent = None
    if (row.get("parent_code") or "").strip():
        parent = location_by_ref(session, tid, row["parent_code"])
        if parent is None:
            raise ValueError(f"unknown parent_code '{row['parent_code']}'")
    place = session.exec(select(Location).where(Location.tenant_id == tid, Location.code == code)).first()
    action = "update" if place else "new"
    if parent is not None and place is not None and parent.location_id == place.location_id:
        raise ValueError("a place can't be its own parent")
    if apply:
        place = place or Location(tenant_id=tid, code=code, name=name, type=ptype)
        place.name, place.type = name, ptype
        place.parent_id = parent.location_id if parent else None
        if row.get("aliases") is not None and row.get("aliases") != "":
            place.alias = "; ".join(_list(row["aliases"]))
        session.add(place)
    return action, f"{ptype} {code} - {name}" + (f" (in {parent.name})" if parent else "")


def _import_employees(session: Session, tid: str, row: dict, *, apply: bool) -> tuple[str, str]:
    email = row.get("email", "").lower()
    if not _EMAIL.match(email):
        raise ValueError(f"invalid email '{email}'")
    name = row.get("name") or ""
    if not name:
        raise ValueError("name is required")
    place = _location_or_error(session, tid, row.get("location") or "")
    people = _people_by_email(session, tid)
    person = people.get(email)
    action = "update" if person else "new"
    manager = None
    mgr_email = (row.get("manager_email") or "").lower().strip()
    if mgr_email:
        manager = people.get(mgr_email)
    if apply:
        if person is None:
            person = Person(tenant_id=tid, email=email, name=name)
            people[email] = person
        person.name = name
        person.department = row.get("department") or person.department
        role = (row.get("role") or person.role or "EMPLOYEE").upper().replace(" ", "_")
        if role in {"COMPANYADMIN", "COMPANY_ADMIN", "ORG_ADMIN"}:
            role = "COMPANY_ADMIN"
        elif role in {"BRANCHADMIN", "BRANCH_ADMIN", "SITE_ADMIN"}:
            role = "BRANCH_ADMIN"
        elif role in {"TEAMOPERATOR", "TEAM_OPERATOR", "TEAM_LEAD"}:
            role = "TEAM_OPERATOR"
        person.role = role
        if manager:
            person.manager_id = manager.person_id
        if place:
            person.site_id = place.location_id
        session.add(person)
    return action, f"{name} <{email}>" + (f" @ {place.name}" if place else "")


def _family_token(name: str) -> str:
    """Map a department name to a short family code (FACILITIES, IT, …)."""
    text = (name or "").upper()
    aliases = (
        ("FACILIT", "FACILITIES"),
        ("MAINT", "FACILITIES"),
        ("HOUSEKEEP", "HOUSEKEEPING"),
        ("TRAVEL", "TRAVEL"),
        ("SECURITY", "SECURITY"),
        ("PROCURE", "PROCUREMENT"),
        ("CAFETERIA", "CAFETERIA"),
        ("CATERING", "CAFETERIA"),
        ("FINANCE", "FINANCE"),
        ("ACCOUNT", "FINANCE"),
        ("HUMAN", "HR"),
        ("PEOPLE", "HR"),
        ("ADMIN", "ADMIN"),
        ("TECH", "IT"),
    )
    if re.search(r"\bHR\b", text):
        return "HR"
    if re.search(r"\bIT\b", text):
        return "IT"
    for needle, code in aliases:
        if needle in text:
            return code
    token = re.sub(r"[^A-Z0-9]+", "_", text).strip("_").split("_")[0]
    return (token or "TEAM")[:20]


def _autocode_department(name: str, place: Optional[Location]) -> str:
    family = _family_token(name)
    if place is None:
        return family
    loc = re.sub(r"[^A-Z0-9]+", "_", (place.alias or place.name or "").upper()).strip("_")
    loc = loc[:24] or place.location_id[:8].upper()
    return f"{family}_{loc}"


def _import_departments(session: Session, tid: str, row: dict, *, apply: bool) -> tuple[str, str]:
    if not row.get("name"):
        raise ValueError("name is required")
    place = _location_or_error(session, tid, row.get("location") or "")
    code = re.sub(r"[^A-Z0-9_]+", "_", (row.get("code") or "").upper()).strip("_")
    if not code:
        code = _autocode_department(row["name"], place)
    for field in ("primary_email", "approver_email"):
        raw = (row.get(field) or "").strip().lower()
        if not raw:
            continue
        if "+" in raw.split("@", 1)[0]:
            raise ValueError(
                f"{field} must be a normal mailbox (no Gmail +tag aliases) — got '{row[field]}'"
            )
        if not _EMAIL.match(raw):
            raise ValueError(f"invalid {field} '{row[field]}'")
    depts = _depts_by_code(session, tid)
    dept = depts.get(code)
    action = "update" if dept else "new"
    if apply:
        if dept is None:
            dept = Department(tenant_id=tid, code=code, name=row["name"])
            depts[code] = dept
        dept.name = row["name"]
        if row.get("categories"):
            dept.categories = [c.lower().replace(" ", "_") for c in _list(row["categories"])]
        if row.get("primary_email"):
            dept.primary_email = row["primary_email"].lower()
        if row.get("backup_emails"):
            dept.backup_emails = [e.lower() for e in _list(row["backup_emails"]) if _EMAIL.match(e.lower())]
        if row.get("approver_email"):
            dept.approver_email = row["approver_email"].lower()
        if row.get("sla_hours"):
            dept.sla_hours = int(_num(row["sla_hours"], 24))
        if row.get("spend_approval_limit"):
            dept.spend_approval_limit = _num(row["spend_approval_limit"])
        if place:
            dept.location_id = place.location_id
        session.add(dept)
    return action, f"{code} - {row['name']}" + (f" @ {place.name}" if place else "")


def _import_stock(session: Session, tid: str, row: dict, *, apply: bool) -> tuple[str, str]:
    code = re.sub(r"[^A-Z0-9_]+", "_", (row.get("sku_code") or row.get("code") or "").upper()).strip("_")
    name = (row.get("name") or "").strip()
    if not code or not name:
        raise ValueError("sku_code and name are required")
    place = _location_or_error(session, tid, row.get("location") or "")
    vendor_email = (row.get("vendor_email") or "").strip().lower()
    vendor = None
    if vendor_email:
        if not _EMAIL.match(vendor_email):
            raise ValueError(f"invalid vendor_email '{row.get('vendor_email')}'")
        vendor = session.exec(
            select(Vendor).where(Vendor.tenant_id == tid, Vendor.contact_email == vendor_email)
        ).first()
    loc_id = place.location_id if place else None
    q = select(StockSku).where(StockSku.tenant_id == tid, StockSku.sku_code == code)
    if loc_id:
        existing = session.exec(q.where(StockSku.location_id == loc_id)).first()
    else:
        existing = session.exec(q.where(StockSku.location_id.is_(None))).first()  # type: ignore[attr-defined]
        if existing is None:
            existing = session.exec(q).first()
    action = "update" if existing else "new"
    if apply:
        sku = existing or StockSku(tenant_id=tid, sku_code=code, name=name)
        sku.name = name
        sku.sku_code = code
        if row.get("qty_on_hand") not in (None, ""):
            sku.qty_on_hand = _num(row.get("qty_on_hand"), sku.qty_on_hand)
        if row.get("reorder_min") not in (None, ""):
            sku.reorder_min = _num(row.get("reorder_min"), sku.reorder_min)
        if row.get("reorder_qty") not in (None, ""):
            sku.reorder_qty = _num(row.get("reorder_qty"), sku.reorder_qty)
        if row.get("unit"):
            sku.unit = str(row["unit"]).strip()[:40]
        if place:
            sku.location_id = place.location_id
        if vendor:
            sku.vendor_id = vendor.vendor_id
        session.add(sku)
    return action, f"{code} - {name}" + (f" @ {place.name}" if place else "")


_RESOURCE_TYPES = {
    "MEETING_ROOM", "PARKING_SLOT", "DESK", "CABIN", "LOCKER",
    "TRAINING_ROOM", "PROJECTOR", "PANTRY", "FIRST_AID_STATION", "SERVER_ROOM",
}
_RESOURCE_ALIASES = {
    "MEETINGROOM": "MEETING_ROOM",
    "PARKING": "PARKING_SLOT",
    "PARKING_SLOT": "PARKING_SLOT",
    "TRAININGROOM": "TRAINING_ROOM",
    "FIRSTAID": "FIRST_AID_STATION",
    "FIRST_AID": "FIRST_AID_STATION",
    "SERVERROOM": "SERVER_ROOM",
}


def _import_resources(session: Session, tid: str, row: dict, *, apply: bool) -> tuple[str, str]:
    name = row.get("name") or ""
    raw_type = (row.get("type") or "MEETING_ROOM").upper().replace(" ", "_").replace("-", "_")
    rtype = _RESOURCE_ALIASES.get(raw_type.replace("_", ""), raw_type)
    rtype = _RESOURCE_ALIASES.get(rtype, rtype)
    if not name:
        raise ValueError("name is required")
    if rtype not in _RESOURCE_TYPES:
        raise ValueError(
            f"type must be one of {', '.join(sorted(_RESOURCE_TYPES))} (got '{row.get('type')}')"
        )
    by_name = _resources_by_name(session, tid)
    res = by_name.get(name.lower())
    action = "update" if res else "new"
    if apply:
        if res is None:
            res = Resource(tenant_id=tid, type=rtype, name=name)
            by_name[name.lower()] = res
        res.type = rtype
        attrs = dict(res.attributes or {})
        if row.get("capacity"):
            attrs["capacity"] = int(_num(row["capacity"]))
        if row.get("floor"):
            attrs["floor"] = row["floor"]
        if row.get("video_conferencing"):
            attrs["video_conferencing"] = _bool(row["video_conferencing"])
        if row.get("display"):
            attrs["display"] = attrs["presentation_display"] = _bool(row["display"])
        if row.get("near_department"):
            attrs["near_department"] = row["near_department"]
        res.attributes = attrs
        session.add(res)
    return action, f"{rtype} {name}"


def _import_vendors(session: Session, tid: str, row: dict, *, apply: bool) -> tuple[str, str]:
    name = row.get("name") or ""
    email = (row.get("contact_email") or "").lower()
    if not name or not _EMAIL.match(email):
        raise ValueError("name and a valid contact_email are required")
    vendor = session.exec(select(Vendor).where(Vendor.tenant_id == tid, Vendor.name == name)).first()
    action = "update" if vendor else "new"
    if apply:
        vendor = vendor or Vendor(tenant_id=tid, name=name, category=row.get("category") or "general", contact_email=email)
        vendor.category = row.get("category") or vendor.category
        vendor.contact_email = email
        vendor.location = row.get("location") or vendor.location
        session.add(vendor)
    return action, name


def _import_knowledge(session: Session, tid: str, row: dict, *, apply: bool) -> tuple[str, str]:
    key = _key(row.get("key") or row.get("title"))
    if not key or not row.get("content"):
        raise ValueError("key (or title) and content are required")
    entry = session.exec(select(KnowledgeEntry).where(KnowledgeEntry.tenant_id == tid, KnowledgeEntry.key == key)).first()
    action = "update" if entry else "new"
    if apply:
        entry = entry or KnowledgeEntry(tenant_id=tid, key=key, title=row.get("title") or key, content=row["content"])
        entry.title = row.get("title") or entry.title
        entry.section = row.get("section") or entry.section
        entry.content = row["content"]
        session.add(entry)
    return action, row.get("title") or key


_DELIVERY_MODELS = {"INCLUDED", "CONTRACT", "SUBSIDISED", "CHARGEABLE", "OUTSOURCED"}
_DELIVERY_ALIASES = {
    "INTERNAL": "INCLUDED",
    "IN_HOUSE": "INCLUDED",
    "INHOUSE": "INCLUDED",
    "VENDOR": "OUTSOURCED",
    "EXTERNAL": "OUTSOURCED",
}


def _import_site_services(session: Session, tid: str, row: dict, *, apply: bool) -> tuple[str, str]:
    code = _key(row.get("code") or row.get("name"))
    site = (row.get("site") or "Corporate Office").strip()
    model = (row.get("delivery_model") or "INCLUDED").strip().upper().replace(" ", "_").replace("-", "_")
    model = _DELIVERY_ALIASES.get(model, model)
    if not code or not row.get("name"):
        raise ValueError("code and name are required")
    if model not in _DELIVERY_MODELS:
        raise ValueError(f"delivery_model must be one of {', '.join(sorted(_DELIVERY_MODELS))} (got '{row.get('delivery_model')}')")
    svc = session.exec(
        select(SiteService).where(SiteService.tenant_id == tid, SiteService.site == site, SiteService.code == code)
    ).first()
    action = "update" if svc else "new"
    if apply:
        svc = svc or SiteService(tenant_id=tid, site=site, code=code, name=row["name"])
        svc.name = row["name"]
        svc.delivery_model = model
        svc.owner_department = (row.get("owner_department") or svc.owner_department or "ADMIN").upper()
        svc.unit = row.get("unit") or svc.unit
        if row.get("rate"):
            svc.rate = _num(row["rate"])
        if row.get("included_limit"):
            svc.included_limit = _num(row["included_limit"])
        if row.get("premium_triggers"):
            svc.premium_triggers = [t.lower() for t in _list(row["premium_triggers"])]
        if row.get("premium_rate"):
            svc.premium_rate = _num(row["premium_rate"])
        if row.get("provider"):
            svc.provider = row["provider"]
        session.add(svc)
    return action, f"{site}: {code} ({model})"


def _import_cost_centres(session: Session, tid: str, row: dict, *, apply: bool) -> tuple[str, str]:
    code = (row.get("code") or "").strip().upper()
    if not code or not row.get("name"):
        raise ValueError("code and name are required")
    email = (row.get("approver_email") or "").strip().lower()
    if email and not _EMAIL.match(email):
        raise ValueError(f"invalid approver_email '{email}'")
    cc = session.exec(select(CostCentre).where(CostCentre.tenant_id == tid, CostCentre.code == code)).first()
    action = "update" if cc else "new"
    if apply:
        cc = cc or CostCentre(tenant_id=tid, code=code, name=row["name"])
        cc.name = row["name"]
        cc.department = row.get("department") or cc.department
        cc.approver_name = row.get("approver_name") or cc.approver_name
        if email:
            cc.approver_email = email
        session.add(cc)
    return action, f"{code} - {row['name']}"
