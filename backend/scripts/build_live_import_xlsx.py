"""Build company_import.xlsx for Setup → Import (one file) using our real test Gmails."""

from pathlib import Path

from openpyxl import Workbook

ADMIN = "skcbsa1221@gmail.com"  # Pune admin + company admin ops
FAC = "flexxftw12@gmail.com"  # BLR admin + facilities base
IT = "it.servicess8989@gmail.com"
HR = "kshitiz.vishwakarma2003@gmail.com"  # BLR employee
CAF = "cateringservices4656@gmail.com"
KAPIL = "kmantri1204@gmail.com"


def sub(addr: str, tag: str) -> str:
    local, domain = addr.split("@")
    return f"{local}+{tag}@{domain}"


def main() -> Path:
    out = Path(__file__).resolve().parents[2] / "company_import.xlsx"
    wb = Workbook()

    ws = wb.active
    ws.title = "locations"
    ws.append(["code", "name", "type", "parent_code", "aliases"])
    for row in [
        ("WEST", "West", "REGION", "", ""),
        ("SOUTH", "South", "REGION", "", ""),
        ("NORTH", "North", "REGION", "", ""),
        ("PUN", "Pune", "CITY", "WEST", ""),
        ("BLR", "Bengaluru", "CITY", "SOUTH", "Bangalore"),
        ("MUM", "Mumbai", "CITY", "WEST", "Bombay"),
        ("GGN", "Gurugram", "CITY", "NORTH", "Gurgaon"),
        ("PUN-HJ", "Pune Hinjewadi", "BRANCH", "PUN", "Hinjewadi; Pune office"),
        ("BLR-HQ", "Bengaluru HQ", "BRANCH", "BLR", "HQ; head office; Bengaluru office; Bangalore HQ"),
        ("MUM-BKC", "Mumbai BKC", "BRANCH", "MUM", "BKC; Mumbai office"),
        ("GGN-CC", "Gurugram Cyber City", "BRANCH", "GGN", "Cyber City; Gurgaon office"),
    ]:
        ws.append(list(row))

    ws = wb.create_sheet("departments")
    ws.append(
        [
            "code",
            "name",
            "categories",
            "primary_email",
            "backup_emails",
            "approver_email",
            "sla_hours",
            "spend_approval_limit",
            "location",
        ]
    )
    fac = "maintenance; furniture; electrical; plumbing; hvac; pest_control; housekeeping"
    it = "it_support; laptop; network; software; access_card; phone_sim; asset"
    admin_cats = "general; meeting_room; supplies; courier; printing; seating; event"
    for row in [
        ("ADMIN", "Admin & Workplace", admin_cats, ADMIN, "", ADMIN, "24", "0", ""),
        ("FACILITIES", "Facilities & Maintenance", fac, sub(FAC, "facilities"), "", ADMIN, "8", "0", ""),
        ("IT", "IT Support", it, sub(IT, "it"), "", ADMIN, "8", "0", ""),
        ("HR", "People / HR", "onboarding; hr_query; offboarding", sub(HR, "hr"), "", ADMIN, "24", "0", ""),
        ("CAFETERIA", "Cafeteria & Catering", "catering; cafeteria", sub(CAF, "cafeteria"), "", ADMIN, "12", "0", ""),
        (
            "SECURITY",
            "Security & Front Desk",
            "visitor; security; parking; material_gate_pass; lost_found; keys_locker",
            sub(ADMIN, "security"),
            "",
            ADMIN,
            "4",
            "0",
            "",
        ),
        (
            "ADMIN_PUN",
            "Admin - Pune",
            "supplies; seating; meeting_room; printing; courier",
            ADMIN,
            "",
            ADMIN,
            "8",
            "0",
            "PUN",
        ),
        (
            "ADMIN_BLR",
            "Admin - Bengaluru",
            "supplies; seating; meeting_room; printing; courier",
            FAC,
            "",
            FAC,
            "8",
            "0",
            "BLR",
        ),
        ("FACILITIES_PUN", "Facilities - Pune", fac, sub(FAC, "pune"), "", ADMIN, "8", "0", "PUN"),
        ("IT_NORTH", "IT Support - North", it, sub(IT, "north"), "", ADMIN, "8", "0", "NORTH"),
    ]:
        ws.append(list(row))

    ws = wb.create_sheet("employees")
    ws.append(["email", "name", "department", "role", "manager_email", "location"])
    for row in [
        (HR, "Kshitiz Vishwakarma", "People", "EMPLOYEE", "", "BLR-HQ"),
        (FAC, "Vikram Shetty", "Engineering", "EMPLOYEE", "", "BLR-HQ"),
        (CAF, "Ananya Rao", "Sales", "EMPLOYEE", "", "PUN-HJ"),
        (IT, "Neha Kapoor", "Finance", "EMPLOYEE", "", "GGN-CC"),
        (ADMIN, "Ops Admin", "Operations", "ADMIN", "", "PUN-HJ"),
        (KAPIL, "Kapil Mantri", "Operations", "EMPLOYEE", "", "BLR-HQ"),
    ]:
        ws.append(list(row))

    ws = wb.create_sheet("resources")
    ws.append(["name", "type", "capacity", "floor", "video_conferencing", "display", "near_department"])
    ws.append(["Board Room BLR-1", "MEETING_ROOM", "12", "3", "yes", "yes", "Sales"])
    ws.append(["Focus Room BLR-2", "MEETING_ROOM", "6", "3", "yes", "no", ""])
    ws.append(["Conference Mumbai BKC", "MEETING_ROOM", "16", "2", "yes", "yes", ""])

    for name, cols in [
        ("vendors", ["name", "category", "contact_email", "location"]),
        ("knowledge", ["key", "section", "title", "content"]),
        (
            "site_services",
            [
                "site",
                "code",
                "name",
                "delivery_model",
                "owner_department",
                "unit",
                "rate",
                "included_limit",
                "premium_triggers",
                "premium_rate",
                "provider",
            ],
        ),
        ("cost_centres", ["code", "name", "department", "approver_name", "approver_email"]),
    ]:
        w = wb.create_sheet(name)
        w.append(cols)

    wb.save(out)
    print(out)
    return out


if __name__ == "__main__":
    main()
