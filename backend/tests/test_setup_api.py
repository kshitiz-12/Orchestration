import io
import json

from openpyxl import Workbook
from sqlmodel import select

from app.agent.directory import department_emails, route_for
from app.models.company import Department

BASE = "/api/v1/setup"


def test_status_shows_env_admin_and_placeholder_departments(client, auth_headers):
    resp = client.get(f"{BASE}/status", headers=auth_headers)
    assert resp.status_code == 200, resp.text
    body = resp.json()
    assert body["admin_email_source"].startswith("env")
    assert body["departments"] >= 10
    assert "FACILITIES" in body["departments_needing_setup"]
    assert body["knowledge_entries"] >= 0


def test_setup_requires_login(client):
    assert client.get(f"{BASE}/departments").status_code == 401


def test_setting_department_email_makes_it_the_real_route(client, auth_headers, engine):
    depts = client.get(f"{BASE}/departments", headers=auth_headers).json()
    fac = next(d for d in depts if d.get("code") == "FACILITIES")
    assert fac["needs_setup"] is True

    resp = client.patch(
        f"{BASE}/departments/{fac['department_id']}",
        json={"primary_email": "Facilities@AcmeCorp.com", "sla_hours": 6},
        headers=auth_headers,
    )
    assert resp.status_code == 200, resp.text
    assert resp.json()["needs_setup"] is False
    assert resp.json()["primary_email"] == "facilities@acmecorp.com"

    from sqlmodel import Session

    with Session(engine) as s:
        tid = s.exec(select(Department)).first().tenant_id
        route = route_for(s, tid, "hvac")
        assert route.recipients == ["facilities@acmecorp.com"]
        assert route.redirected_to_admin is False
        assert "facilities@acmecorp.com" in department_emails(s, tid)


def test_create_and_deactivate_department(client, auth_headers):
    resp = client.post(
        f"{BASE}/departments",
        json={"code": "legal", "name": "Legal", "categories": ["Contract Review"], "primary_email": "legal@acmecorp.com"},
        headers=auth_headers,
    )
    assert resp.status_code == 200, resp.text
    dept = resp.json()
    assert dept["code"] == "LEGAL" and dept["categories"] == ["contract_review"]
    assert client.post(f"{BASE}/departments", json={"code": "LEGAL", "name": "x"}, headers=auth_headers).status_code == 409
    assert client.delete(f"{BASE}/departments/{dept['department_id']}", headers=auth_headers).json()["ok"] is True

    admin = next(d for d in client.get(f"{BASE}/departments", headers=auth_headers).json() if d["code"] == "ADMIN")
    assert client.delete(f"{BASE}/departments/{admin['department_id']}", headers=auth_headers).status_code == 400


def test_bad_department_email_rejected(client, auth_headers):
    resp = client.post(f"{BASE}/departments", json={"code": "X", "name": "X", "primary_email": "not an email"}, headers=auth_headers)
    assert resp.status_code == 422


def test_knowledge_upsert_and_delete(client, auth_headers):
    resp = client.put(
        f"{BASE}/knowledge",
        json={"key": "Canteen Hours", "title": "Canteen hours", "content": "8:30 AM to 6 PM", "section": "faq"},
        headers=auth_headers,
    )
    assert resp.json()["key"] == "canteen_hours"
    client.put(f"{BASE}/knowledge", json={"key": "canteen_hours", "title": "Canteen hours", "content": "9 AM to 7 PM"}, headers=auth_headers)
    rows = [r for r in client.get(f"{BASE}/knowledge", headers=auth_headers).json() if r["key"] == "canteen_hours"]
    assert len(rows) == 1 and rows[0]["content"] == "9 AM to 7 PM"
    assert client.delete(f"{BASE}/knowledge/{rows[0]['entry_id']}", headers=auth_headers).json()["ok"] is True


def test_template_download(client, auth_headers):
    resp = client.get(f"{BASE}/templates/employees", headers=auth_headers)
    assert resp.status_code == 200
    assert resp.text.splitlines()[0] == "email,name,department,role,manager_email,location"
    assert client.get(f"{BASE}/templates/nope", headers=auth_headers).status_code == 404


def test_employee_csv_preview_then_commit(client, auth_headers):
    csv_text = (
        "Email,Name,Department,Role\n"
        "riya.shah@acmecorp.com,Riya Shah,Engineering,employee\n"
        "bad-email,Nobody,,\n"
        "admin@prototype.local,Prototype Admin,Admin,ADMIN\n"
    )
    files = {"file": ("people.csv", csv_text.encode(), "text/csv")}
    preview = client.post(f"{BASE}/import/employees", files=files, headers=auth_headers).json()
    assert preview["committed"] is False
    assert preview["new"] == 1 and preview["updated"] == 1
    assert preview["errors"][0]["row"] == 3
    before = client.get(f"{BASE}/status", headers=auth_headers).json()["employees"]

    files = {"file": ("people.csv", csv_text.encode(), "text/csv")}
    done = client.post(f"{BASE}/import/employees?commit=true", files=files, headers=auth_headers).json()
    assert done["committed"] is True and done["new"] == 1
    assert client.get(f"{BASE}/status", headers=auth_headers).json()["employees"] == before + 1


def test_department_xlsx_import_updates_existing(client, auth_headers):
    wb = Workbook()
    ws = wb.active
    ws.append(["code", "name", "categories", "primary_email"])
    ws.append(["IT", "IT Support", "it_support; laptop; vpn", "it@acmecorp.com"])
    ws.append(["LEGAL", "Legal", "contract_review", "legal@acmecorp.com"])
    buf = io.BytesIO()
    wb.save(buf)
    files = {"file": ("depts.xlsx", buf.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    result = client.post(f"{BASE}/import/departments?commit=true", files=files, headers=auth_headers).json()
    assert result["updated"] == 1 and result["new"] == 1 and not result["errors"]
    it = next(d for d in client.get(f"{BASE}/departments", headers=auth_headers).json() if d["code"] == "IT")
    assert it["primary_email"] == "it@acmecorp.com" and "vpn" in it["categories"] and it["needs_setup"] is False


def test_empty_upload_rejected(client, auth_headers):
    files = {"file": ("empty.csv", b"email,name\n", "text/csv")}
    assert client.post(f"{BASE}/import/employees", files=files, headers=auth_headers).status_code == 422


def test_company_workbook_preview_and_commit(client, auth_headers):
    wb = Workbook()
    loc = wb.active
    loc.title = "locations"
    loc.append(["code", "name", "type", "parent_code", "aliases"])
    loc.append(["WEST", "West", "REGION", "", ""])
    loc.append(["PUN", "Pune", "CITY", "WEST", ""])
    loc.append(["PUN-HJ", "Pune Hinjewadi", "BRANCH", "PUN", "Hinjewadi"])
    dept = wb.create_sheet("departments")
    dept.append(["code", "name", "categories", "primary_email", "backup_emails", "approver_email", "sla_hours",
                 "spend_approval_limit", "location"])
    dept.append(["ADMIN_PUN", "Admin - Pune", "supplies; seating", "skcbsa1221@gmail.com", "",
                 "skcbsa1221@gmail.com", "8", "0", "PUN"])
    emp = wb.create_sheet("employees")
    emp.append(["email", "name", "department", "role", "manager_email", "location"])
    emp.append(["riya.shah@acmecorp.com", "Riya Shah", "Engineering", "EMPLOYEE", "", "PUN-HJ"])
    buf = io.BytesIO()
    wb.save(buf)
    files = {"file": ("company.xlsx", buf.getvalue(),
                      "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}

    preview = client.post(f"{BASE}/import/company", files=files, headers=auth_headers).json()
    assert preview["committed"] is False
    assert preview["source"] == "workbook"
    assert set(preview["kinds"]) == {"locations", "departments", "employees"}
    assert preview["new"] >= 4
    before = client.get(f"{BASE}/status", headers=auth_headers).json()["employees"]

    files = {"file": ("company.xlsx", buf.getvalue(),
                      "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    done = client.post(f"{BASE}/import/company?commit=true", files=files, headers=auth_headers).json()
    assert done["committed"] is True
    assert client.get(f"{BASE}/status", headers=auth_headers).json()["employees"] == before + 1
    admin_pun = next(
        d for d in client.get(f"{BASE}/departments", headers=auth_headers).json() if d["code"] == "ADMIN_PUN"
    )
    assert admin_pun["primary_email"] == "skcbsa1221@gmail.com"


def test_company_ai_json_dump_import(client, auth_headers):
    payload = {
        "locations": [
            {"code": "BLR", "name": "Bengaluru", "type": "CITY", "parent_code": "", "aliases": "Bangalore"},
            {"code": "BLR-HQ", "name": "Bengaluru HQ", "type": "BRANCH", "parent_code": "BLR", "aliases": "HQ"},
        ],
        "departments": [
            {
                "code": "ADMIN_BLR",
                "name": "Admin - Bengaluru",
                "categories": "meeting_room; supplies",
                "primary_email": "flexxftw12@gmail.com",
                "backup_emails": "",
                "approver_email": "flexxftw12@gmail.com",
                "sla_hours": "8",
                "spend_approval_limit": "0",
                "location": "BLR",
            }
        ],
        "employees": [
            {
                "email": "kshitiz.vishwakarma2003@gmail.com",
                "name": "Kshitiz Vishwakarma",
                "department": "People",
                "role": "EMPLOYEE",
                "manager_email": "",
                "location": "BLR-HQ",
            }
        ],
    }
    files = {"file": ("org.json", json.dumps(payload).encode(), "application/json")}
    done = client.post(f"{BASE}/import/company?commit=true", files=files, headers=auth_headers).json()
    assert done["committed"] is True
    assert done["source"] == "ai"
    assert "ADMIN_BLR" in {d["code"] for d in client.get(f"{BASE}/departments", headers=auth_headers).json()}


def test_company_template_download(client, auth_headers):
    resp = client.get(f"{BASE}/templates/company", headers=auth_headers)
    assert resp.status_code == 200
    assert "spreadsheetml" in resp.headers.get("content-type", "")
